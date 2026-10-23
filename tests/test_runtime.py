"""Synthetic end-to-end tests of the production orchestration interfaces.
No real API/model invocation; all customer identities are reserved .example.
"""
import json
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from openpyxl import load_workbook
from yingba.runtime import Runtime, settings
from yingba.db import Store
from yingba.engine import Engine
from yingba.budget import Budget, BudgetPause, beijing_day
from yingba.paid import PaidGateway
from yingba.adapters import ConnectionPause
from yingba.apollo import parse_search_response
from yingba.public_web import PublicVerifier, Page
from yingba.exchange import export

class Client:
    calls=[]
    def __init__(self, api_key, enabled=False):
        self.key=api_key
        if api_key!='synthetic-fixture-key':raise ValueError('invalid fixture')
    def health(self):return {'healthy':True,'is_logged_in':True,'search_permission_verified':False}
    def search(self, request, authorized=False):
        assert authorized
        self.calls.append(request)
        p=request['page']
        rows=[{'id':str(i),'name':f'Fixture Oil {i}','website_url':f'https://oil{i}.example'} for i in range((p-1)*2+1,p*2+1)]
        return parse_search_response({'organizations':rows,'pagination':{'page':p,'per_page':2,'total_entries':4,'total_pages':2}},request)
class Fetcher:
    def fetch(self,url,expected_domain):
        n=expected_domain.removeprefix('oil').split('.')[0]
        return Page(url,f'<title>Fixture Oil {n}</title><h1>Fixture Oil {n}</h1><p>We are a lubricant distributor.</p><p>Contact sales@{expected_domain}</p>')
class Reviewer:
    def preflight(self):return {'connected':True}
    def review_company(self,record,evidence,enabled=False):
        assert enabled
        return {'identity_confirmed':True,'business_confirmed':True,'category':'润滑油经销商',
                'identity_refs':[{'evidence_id':'0','quote':record['name']}],
                'business_refs':[{'evidence_id':'0','quote':'We are a lubricant distributor.'}],
                'email_refs':[{'email':'sales@'+record['domain'],'evidence_id':'0','quote':'sales@'+record['domain']}], 'reason':'合成接口测试'}

class RuntimeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.store=Store(self.root/'db');Client.calls=[]
        self.rt=Runtime(self.store,client_factory=Client,verifier=PublicVerifier(Fetcher()),reviewer=Reviewer())
        self.engine=Engine(self.store,source_factory=self.rt.source_for_task)
        self.config=settings({'estimate_points':2,'per_page':2})
    def tearDown(self):self.temp.cleanup()
    def create(self,target=3):return self.engine.create(target,'real',config=self.config)
    def connect(self):self.rt.connect('synthetic-fixture-key',self.config['account_label'],True)
    def settle(self,action):
        PaidGateway(self.store).reconcile_captured(action,1,beijing_day(),action+'-bill')
    def test_real_adapter_pipeline_refill_surplus_recovery_and_export(self):
        t=self.create();self.connect();Budget(self.store).confirm(0)
        self.assertFalse(self.engine.step(t))
        pending=PaidGateway(self.store).pending();self.assertEqual(pending[0]['response_saved'],1)
        self.assertEqual(self.engine.counts(t)['合格有邮箱'],0)
        self.assertEqual(len(Client.calls),1)
        self.settle(pending[0]['action'])
        # Restart loses key but durable reconciled search response can be replayed.
        fresh=Runtime(Store(self.store.path),client_factory=Client,verifier=PublicVerifier(Fetcher()),reviewer=Reviewer())
        e=Engine(fresh.store,source_factory=fresh.source_for_task);e.resume(t)
        self.assertTrue(e.step(t));self.assertEqual(len(Client.calls),1)
        e.step(t);e.step(t)
        self.assertEqual(e.counts(t)['合格有邮箱'],2)
        self.assertFalse(e.step(t));self.assertEqual(len(Client.calls),1)
        fresh.connect('synthetic-fixture-key',self.config['account_label'],True)
        Budget(self.store).confirm(1);e.resume(t);e.step(t)
        self.assertEqual(len(Client.calls),2)
        self.settle(PaidGateway(self.store).pending()[0]['action'])
        e.resume(t);e.run(t)
        self.assertEqual(e.counts(t)['合格有邮箱'],4)
        self.assertEqual(self.store.rows('SELECT status FROM tasks')[0]['status'],'complete')
        p,_=export(self.store,t,self.root/'exports');w=load_workbook(p)
        self.assertEqual(w['公司总表'].max_row-1,4);self.assertEqual(w['唯一邮箱清单'].max_row-1,4);w.close()
        self.assertNotIn('synthetic-fixture-key',Path(self.store.path).read_bytes().decode('latin1'))
    def test_no_downstream_acceptance_means_no_paid_request(self):
        self.rt.reviewer.preflight=lambda:{'connected':False,'reason':'需要本人订阅验收'}
        t=self.create();self.connect();Budget(self.store).confirm(0);self.engine.step(t)
        self.assertEqual(Client.calls,[]);self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
    def test_unknown_price_pauses_before_reservation(self):
        self.config['estimate_points']=None;t=self.create();self.connect();Budget(self.store).confirm(0)
        self.engine.step(t);self.assertEqual(Client.calls,[]);self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
    def test_unconfirmed_account_budget_stops_network(self):
        t=self.create();self.connect();self.engine.step(t)
        self.assertEqual(Client.calls,[])
    def test_key_connection_needs_explicit_authorization(self):
        with self.assertRaises(ConnectionPause):self.rt.connect('synthetic-fixture-key','本机唯一Apollo账号',False)
        self.assertFalse(self.rt.status()['health_checked'])
    def test_account_mismatch_stops_network(self):
        t=self.create();self.rt.connect('synthetic-fixture-key','另一个账号',True);Budget(self.store).confirm(0)
        self.engine.step(t);self.assertEqual(Client.calls,[])
    def test_task_settings_immutable_and_restart_keeps_them(self):
        t=self.create()
        with self.assertRaises(Exception):self.rt.save_settings(t,{'filters':{'q_organization_keyword_tags':['other']}})
        got=json.loads(Store(self.store.path).rows('SELECT config FROM task_settings')[0]['config'])
        self.assertEqual(got,self.config)
    def test_response_cannot_auto_settle_or_retry(self):
        t=self.create();self.connect();Budget(self.store).confirm(0);self.engine.step(t)
        self.engine.resume(t);self.engine.step(t)
        self.assertEqual(len(Client.calls),1)
        self.assertEqual(self.store.rows('SELECT status FROM charges')[0]['status'],'unknown')
        self.assertEqual(self.store.rows('SELECT page FROM tasks')[0]['page'],1)
    def test_account_switch_between_reserve_and_send_cannot_use_new_key(self):
        t=self.create();self.connect();Budget(self.store).confirm(0)
        original=PaidGateway.execute
        def switch(gateway,*args):
            self.rt.connect('synthetic-fixture-key','另一个账号',True)
            return original(gateway,*args)
        with patch('yingba.runtime.PaidGateway.execute',switch):self.engine.step(t)
        self.assertEqual(Client.calls,[])
        self.assertEqual(self.store.rows('SELECT status FROM charges')[0]['status'],'unknown')
    def test_midnight_between_reservation_and_send_blocks_network(self):
        t=self.create();self.connect();Budget(self.store).confirm(0)
        with patch('yingba.runtime.beijing_day',return_value='2099-01-01'):self.engine.step(t)
        self.assertEqual(Client.calls,[])
    def test_model_cannot_reintroduce_privacy_or_hr_mailboxes(self):
        class Private(Fetcher):
            def fetch(self,url,expected_domain):
                return Page(url,'<title>Fixture Oil 1</title><p>We are a lubricant distributor.</p><p>privacy@oil1.example hr@oil1.example sales@oil1.example</p>')
        class BadReview(Reviewer):
            def review_company(inner,record,evidence,enabled=False):
                result=super().review_company(record,evidence,enabled)
                result['email_refs']=[{'email':a+'@oil1.example','evidence_id':'0','quote':a+'@oil1.example'} for a in ['privacy','hr','sales']]
                return result
        self.rt.verifier=PublicVerifier(Private());self.rt.reviewer=BadReview()
        source=self.rt.source_for_task({'id':self.create(),'mode':'real'})
        result=source.verify({'name':'Fixture Oil 1','domain':'oil1.example'})
        self.assertEqual(result['emails'],['sales@oil1.example'])
    def test_fetched_evidence_survives_reviewer_failure_and_restart(self):
        class Denied(Reviewer):
            def review_company(self,*a,**k):raise ConnectionPause('订阅限额')
        self.rt.reviewer=Denied();t=self.create();record={'name':'Fixture Oil 1','domain':'oil1.example'}
        with self.assertRaises(ConnectionPause):self.rt.source_for_task({'id':t,'mode':'real'}).verify(record)
        self.assertEqual(len(self.store.rows('SELECT * FROM verification_cache')),1)
        self.rt.reviewer=Reviewer()
        with patch.object(self.rt.verifier,'verify',side_effect=AssertionError('不应重抓')):
            result=self.rt.source_for_task({'id':t,'mode':'real'}).verify(record)
        self.assertFalse(result['review_required'])
    def test_rules_cannot_be_overridden_by_semantic_assertions(self):
        class Wrong(Fetcher):
            def fetch(self,url,expected_domain):return Page(url,'<title>Unrelated news</title><p>sales@'+expected_domain+'</p>')
        self.rt.verifier=PublicVerifier(Wrong());source=self.rt.source_for_task({'id':self.create(),'mode':'real'})
        result=source.verify({'name':'Fixture Oil 1','domain':'oil1.example'})
        self.assertFalse(result['identity_confirmed']);self.assertFalse(result['business_confirmed'])

if __name__=='__main__':unittest.main()
