"""Existing Apollo-file import tests; no service, browser or model calls."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook

from yingba.apollo_import import prepare_import, page_import
from yingba.db import Store
from yingba.engine import Engine


SOURCE = 'https://app.apollo.io/#/companies?finderViewId=fixture'
CSV = (b'Company name,Website,Apollo account id,Apollo contact id,Email,Email status\r\n'
       b'Example Oils,https://www.oils.example,A-1,C-1,sales@oils.example,verified\r\n'
       b'Example Oils,https://oils.example,A-1,C-2,other@oils.example,verified\r\n'
       b'Other Oils,other.example,A-2,C-3,privacy@other.example,verified\r\n')


class ApolloImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / 'import.sqlite')

    def tearDown(self):
        self.temp.cleanup()

    def upload(self, content=CSV, filename='apollo.csv', source=SOURCE, mapping=None):
        return prepare_import(self.store, content, filename, source, True, mapping)

    def xlsx(self, rows, simulation=False):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        if simulation:
            sheet = book.create_sheet('导出说明')
            sheet.append(['模式', '模拟数据，禁止作为真实客户使用'])
        out = BytesIO()
        book.save(out)
        book.close()
        return out.getvalue()

    def test_original_bytes_hash_raw_rows_and_untrusted_provenance_persist(self):
        result = self.upload()
        self.assertEqual(result['rows'], 3)
        batch = self.store.rows('SELECT * FROM apollo_imports')[0]
        self.assertEqual(batch['content'], CSV)
        self.assertEqual(batch['source_url'], SOURCE)
        self.assertIn('未在线认证', batch['source_attestation'])
        raw = json.loads(self.store.rows('SELECT raw FROM apollo_import_rows ORDER BY row_number')[0]['raw'])
        self.assertEqual(raw['values'][-2:], ['sales@oils.example', 'verified'])
        self.assertEqual(self.store.rows('SELECT COUNT(*) n FROM charges')[0]['n'], 0)

    def test_people_export_never_uses_contact_record_id_as_organization(self):
        result = self.upload(CSV.replace(b'Apollo contact id', b'Apollo record id'))
        records, exhausted = page_import(self.store, result['batch_id'], 1, 100)
        self.assertTrue(exhausted)
        self.assertEqual(records[0]['apollo_id'], '')
        self.assertEqual(records[0]['apollo_account_id'], 'A-1')
        self.assertEqual(records[0]['domain'], 'oils.example')
        self.assertNotIn('emails', records[0])
        self.assertNotIn('verified', records[0])

    def test_only_explicit_organization_column_becomes_organization_id(self):
        data = b'Account Name,Account Website (Domain),Organization ID,Apollo account id\nFixture,fixture.example,O-1,A-1\n'
        result = self.upload(data)
        record = page_import(self.store, result['batch_id'], 1, 100)[0][0]
        self.assertEqual(record['apollo_id'], 'O-1')
        self.assertEqual(record['apollo_account_id'], 'A-1')

    def test_explicit_mapping_cannot_promote_contact_id(self):
        with self.assertRaisesRegex(ValueError, '联系人'):
            self.upload(mapping={'organization_id': 'Apollo contact id'})

    def test_accounts_record_id_not_assumed_to_be_organization_id(self):
        result = self.upload(b'Account Name,Account Website (Domain),Apollo record ID\nFixture,fixture.example,A-1\n')
        record = page_import(self.store, result['batch_id'], 1, 100)[0][0]
        self.assertEqual(record['apollo_id'], '')
        self.assertEqual(record['apollo_account_id'], '')

    def test_paging_is_stable_across_restart_and_exhaustion(self):
        batch = self.upload()['batch_id']
        one, exhausted = page_import(self.store, batch, 1, 2)
        self.assertEqual(len(one), 2)
        self.assertFalse(exhausted)
        two, exhausted = page_import(Store(self.store.path), batch, 2, 2)
        self.assertEqual(len(two), 1)
        self.assertTrue(exhausted)
        self.assertEqual(page_import(self.store, batch, 3, 2), ([], True))
        self.assertEqual(one[0]['apollo_import']['row_number'], 2)

    def test_duplicate_import_is_idempotent_even_concurrent(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: self.upload(), range(4)))
        self.assertEqual(len({row['batch_id'] for row in results}), 1)
        self.assertEqual(sum(not row['duplicate'] for row in results), 1)
        self.assertEqual(len(self.store.rows('SELECT * FROM apollo_imports')), 1)
        self.assertEqual(len(self.store.rows('SELECT * FROM apollo_import_rows')), 3)

    def test_source_and_mapping_are_bound_to_batch(self):
        data = b'Company name,Account Name,Website\nName One,Name Two,fixture.example\n'
        first = self.upload(data, mapping={'name': 'Company name'})
        second = self.upload(data, mapping={'name': 'Account Name'})
        third = self.upload(data, source=SOURCE + '-new', mapping={'name': 'Account Name'})
        self.assertEqual(len({r['batch_id'] for r in (first, second, third)}), 3)
        self.assertTrue(self.upload(data, mapping={'name': 'Company name'})['duplicate'])

    def test_ambiguous_columns_require_mapping(self):
        with self.assertRaisesRegex(ValueError, '歧义'):
            self.upload(b'Company name,Account Name,Website\nA,B,x.example\n')

    def test_duplicate_case_whitespace_column_names_rejected(self):
        with self.assertRaisesRegex(ValueError, '重名'):
            self.upload(b'Company name, Company Name ,Website\nA,B,x.example\n')

    def test_utf8_bom_utf16_and_chinese_paths(self):
        text = 'Company name,Website\n润滑公司,oil.example\n'
        for encoding in ('utf-8-sig', 'utf-16', 'utf-16-be'):
            content = text.encode(encoding)
            if encoding == 'utf-16-be':
                content = b'\xfe\xff' + content
            result = self.upload(content, filename='中文目录\\已有 Apollo 名单.csv')
            record = page_import(self.store, result['batch_id'], 1, 100)[0][0]
            self.assertEqual(record['name'], '润滑公司')

    def test_unsupported_encoding_and_null_bytes_rejected(self):
        for content in ('Company name,Website\n润滑公司,x.example\n'.encode('gbk'), b'C\x00o\x00m\x00p\x00a\x00n\x00y\x00'):
            with self.assertRaises(ValueError):
                self.upload(content)

    def test_xlsx_import_preserves_rows_and_original(self):
        content = self.xlsx([['Company name', 'Website', 'Email'], ['Fixture', 'fixture.example', 'sales@fixture.example']])
        result = self.upload(content, 'company.xlsx')
        self.assertEqual(self.store.rows('SELECT content FROM apollo_imports')[0]['content'], content)
        record = page_import(self.store, result['batch_id'], 1, 100)[0][0]
        self.assertNotIn('emails', record)

    def test_missing_name_or_domain_rows_preserved_for_engine_exclusion(self):
        result = self.upload(b'Company name,Website\n,fixture.example\nFixture,\n')
        records = page_import(self.store, result['batch_id'], 1, 100)[0]
        self.assertEqual(records[0]['name'], '')
        self.assertEqual(records[1]['domain'], '')
        engine = Engine(self.store)
        task = engine.create(1000, 'real')
        with self.store.transaction() as db:
            for record in records:
                engine._ingest(db, task, 1, record)
        self.assertEqual(engine.counts(task)['合格有邮箱'], 0)
        self.assertEqual(len(self.store.rows('SELECT * FROM raw_records')), 2)

    def test_people_rows_are_company_deduplicated_not_email_qualified(self):
        result = self.upload()
        records = page_import(self.store, result['batch_id'], 1, 100)[0]
        engine = Engine(self.store)
        task = engine.create(1000, 'real')
        with self.store.transaction() as db:
            for record in records:
                engine._ingest(db, task, 1, record)
        self.assertEqual(engine.counts(task)['候选'], 2)
        self.assertEqual(engine.counts(task)['合格有邮箱'], 0)
        self.assertEqual(self.store.rows('SELECT * FROM emails'), [])
        self.assertEqual(len(self.store.rows('SELECT * FROM raw_records')), 3)

    def test_no_website_is_guessed_from_email_or_name(self):
        with self.assertRaisesRegex(ValueError, '不从'):
            self.upload(b'Company name,Email\nCompany,sales@company.example\n')

    def test_simulation_exports_filename_metadata_and_column_rejected(self):
        with self.assertRaisesRegex(ValueError, '模拟'):
            self.upload(filename='SIMULATION_custom.csv')
        with self.assertRaisesRegex(ValueError, '模拟'):
            self.upload(self.xlsx([['Company name', 'Website'], ['A', 'a.example']], simulation=True), 'a.xlsx')
        with self.assertRaisesRegex(ValueError, '模拟'):
            self.upload(b'Company name,Website,simulation\nA,a.example,true\n')

    def test_sensitive_credentials_columns_never_stored(self):
        for column in ('Apollo API Key', 'Authorization Header', 'Refresh Token', 'Password Hash', 'Cookie', 'OpenAI API Key', 'API Token'):
            with self.assertRaisesRegex(ValueError, '敏感'):
                self.upload(f'Company name,Website,{column}\nA,a.example,synthetic-secret\n'.encode())
        self.assertEqual(self.store.rows('SELECT * FROM apollo_imports'), [])

    def test_hidden_sheet_credentials_rejected_before_original_saved(self):
        book = Workbook()
        book.active.append(['Company name', 'Website'])
        book.active.append(['Fixture', 'fixture.example'])
        hidden = book.create_sheet('private')
        hidden.sheet_state = 'hidden'
        hidden.append(['API Token'])
        hidden.append(['synthetic-secret'])
        out = BytesIO()
        book.save(out)
        book.close()
        with self.assertRaisesRegex(ValueError, '敏感'):
            self.upload(out.getvalue(), 'hidden.xlsx')
        self.assertEqual(self.store.rows('SELECT * FROM apollo_imports'), [])

    def test_explicit_source_attestation_and_official_https_required(self):
        with self.assertRaisesRegex(ValueError, '确认'):
            prepare_import(self.store, CSV, 'a.csv', SOURCE, False)
        for url in ('http://app.apollo.io', 'https://app.apollo.io.evil.example/', 'https://notapollo.io/', 'https://key@app.apollo.io/', 'https://app.apollo.io:8000/', 'https://app.apollo.io/\r\n'):
            with self.assertRaises(ValueError):
                self.upload(source=url)

    def test_harmless_official_urls_accepted_as_declarations_only(self):
        for url in ('https://app.apollo.io/#/people', 'https://apollo.io/', 'https://www.apollo.io/'):
            self.assertIn('未在线认证', self.upload(source=url)['source_attestation'])

    def test_credentials_in_query_or_hash_source_url_never_saved(self):
        for ending in ('?api_key=synthetic', '#access_token=synthetic',
                       '#/people?password=synthetic', '?%61pi%5Fkey=synthetic',
                       '#/companies?finderViewId=fixture&refresh_token=synthetic',
                       '?code=synthetic', '?x=1;token=synthetic'):
            with self.assertRaises(ValueError):
                self.upload(source='https://app.apollo.io/' + ending)
        self.assertEqual(self.store.rows('SELECT * FROM apollo_imports'), [])

    def test_bad_size_type_and_format_rejected(self):
        for content, filename in ((b'', 'a.csv'), ('bad', 'a.csv'), (b'x', 'auth.json'), (b'x', 'a.xlsm'), (b'x', 'a.xls')):
            with self.assertRaises(ValueError):
                self.upload(content, filename)
        with patch('yingba.apollo_import.MAX_BYTES', 10):
            with self.assertRaisesRegex(ValueError, '10MB'):
                self.upload()

    def test_row_limit_and_zip_members_limit_leave_no_partial_batch(self):
        with patch('yingba.apollo_import.MAX_ROWS', 2):
            with self.assertRaisesRegex(ValueError, '10万'):
                self.upload()
        out = BytesIO()
        with zipfile.ZipFile(out, 'w') as archive:
            for i in range(201):
                archive.writestr(f'file{i}', 'x')
        with self.assertRaisesRegex(ValueError, '展开过大'):
            self.upload(out.getvalue(), 'large.xlsx')
        self.assertEqual(self.store.rows('SELECT * FROM apollo_imports'), [])

    def test_malformed_csv_extra_values_and_missing_headers_rejected(self):
        for content in (b'Company name,Website\nA,a.example,unexpected\n', b'Company name,Website\n"unclosed', b'Company name,Website\n', b'First Name,Last Name\nA,B\n'):
            with self.assertRaises(ValueError):
                self.upload(content)

    def test_formula_cells_are_retained_as_untrusted_text_not_executed(self):
        content = self.xlsx([['Company name', 'Website', 'Notes'], ['Fixture', 'fixture.example', '=HYPERLINK("https://example.invalid")']])
        result = self.upload(content, 'formula.xlsx')
        record = page_import(self.store, result['batch_id'], 1, 100)[0][0]
        self.assertTrue(record['apollo_import']['raw']['values'][2].startswith('=HYPERLINK'))

    def test_page_validation_and_unknown_batch(self):
        batch = self.upload()['batch_id']
        for page, size in ((0, 100), (True, 100), (1, 0), (1, 1001), (1, True)):
            with self.assertRaises(ValueError):
                page_import(self.store, batch, page, size)
        with self.assertRaises(ValueError):
            page_import(self.store, '0' * 64, 1, 100)
        with self.assertRaises(ValueError):
            page_import(self.store, '../secret', 1, 100)


if __name__ == '__main__':
    unittest.main()
