"""No paid API, no model: zero-credit and local-import orchestration contracts."""
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from yingba.apollo import parse_accounts_response, ApolloPause
from yingba.apollo_import import prepare_import
from yingba.adapters import ConnectionPause
from yingba.budget import Budget
from yingba.db import Store
from yingba.engine import Engine
from yingba.exchange import export,import_history
from yingba.runtime import Runtime, settings, append_import
from yingba.public_web import PublicVerifier
from tests.test_runtime import Fetcher, Reviewer

class FreeClient:
    calls=[]
    def __init__(self,*a,**k):pass
    def health(self):return {'healthy':True,'is_logged_in':True}
    def search(self,*a,**k):raise AssertionError('Never use paid search or enrichment')
    def search_accounts(self,request):
        self.calls.append(request)
        p=request['page'];n=request['per_page'];total=4
        accounts=[{'id':f'acct-{i}','organization_id':f'org-{i}','name':f'Fixture Oil {i}','domain':f'oil{i}.example'} for i in range((p-1)*n+1,min(p*n,total)+1)]
        return parse_accounts_response({'accounts':accounts,'pagination':{'page':p,'per_page':n,'total_entries':total,'total_pages':(total+n-1)//n}},request)

class FakeCosts:
    """Injected fee decision; no official-site fetch, API or model invocation."""
    def assess(self,**kwargs):
        if kwargs.get('source') != 'apollo_accounts':raise AssertionError('Only the saved-account free fixture is authorized')
        return {'classification':'free','max_points':0,'reason':'Synthetic documented-zero fixture'}

class FreeRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.store=Store(self.root/'db')
        self.rt=Runtime(self.store,client_factory=FreeClient,verifier=PublicVerifier(Fetcher()),reviewer=Reviewer(),cost_service=FakeCosts())
        self.e=Engine(self.store,source_factory=self.rt.source_for_task);FreeClient.calls=[]
    def tearDown(self):self.temp.cleanup()
    def free_task(self,target=3):
        config=settings({'source':'apollo_accounts','per_page':2})
        t=self.e.create(target,'real',config=config)
        self.rt.connect('synthetic','本机唯一Apollo账号',True)
        return t
    def imported(self,ids=(1,2),target=3):
        raw=('Company Name,Website,Email,Apollo Account ID\n'+''.join(f'Fixture Oil {i},oil{i}.example,unverified@oil{i}.example,acct-{i}\n' for i in ids)).encode()
        info=prepare_import(self.store,raw,'apollo.csv','https://app.apollo.io/#/accounts',True)
        t=self.e.create(target,'real',config=settings({'source':'apollo_import','batch_id':info['batch_id'],'per_page':2}))
        return t,info
    def test_documented_free_automatically_refills_without_budget_or_reconciliation(self):
        t=self.free_task()
        with patch('yingba.runtime.PaidGateway.execute',side_effect=AssertionError('paid gateway')):self.e.run(t)
        self.assertEqual(self.e.counts(t)['合格有邮箱'],4)
        self.assertEqual(self.store.rows('SELECT status FROM tasks')[0]['status'],'complete')
        self.assertEqual(len(FreeClient.calls),2)
        self.assertEqual(len(self.store.rows('SELECT * FROM free_pages')),2)
        self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
        self.assertEqual(self.store.rows('SELECT * FROM budgets'),[])
    def test_unknown_paid_ledger_is_untouched_by_free_reads(self):
        Budget(self.store).confirm(0);Budget(self.store).reserve('old-unknown',3);Budget(self.store).unknown('old-unknown')
        before=self.store.rows('SELECT * FROM charges');t=self.free_task();self.e.run(t)
        self.assertEqual(self.store.rows('SELECT * FROM charges'),before)
        self.assertEqual(self.e.counts(t)['合格有邮箱'],4)
    def test_free_source_exhaustion_never_falls_back_to_paid_search(self):
        t=self.free_task(1000);self.e.run(t)
        self.assertEqual(self.store.rows('SELECT status FROM tasks')[0]['status'],'insufficient')
        self.assertEqual(self.e.counts(t)['距目标差额'],996)
        self.assertEqual(len(FreeClient.calls),2)
    def test_permission_denial_preserves_page_without_fee_or_retry(self):
        t=self.free_task()
        with patch.object(self.rt._client,'search_accounts',side_effect=ApolloPause('permission','无读取权限')) as call:
            self.e.run(t);self.e.run(t);self.assertEqual(call.call_count,1)
        row=self.store.rows('SELECT status,page FROM tasks')[0]
        self.assertEqual(row,{'status':'paused','page':1});self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
    def test_cached_free_page_survives_crash_before_page_commit_without_key(self):
        t=self.free_task();source=self.rt.source_for_task({'id':t,'mode':'real'});first=source.search(1)
        fresh=Runtime(Store(self.store.path),client_factory=FreeClient,reviewer=Reviewer(),cost_service=FakeCosts())
        again=fresh.source_for_task({'id':t,'mode':'real'}).search(1)
        self.assertEqual(again,first);self.assertEqual(len(FreeClient.calls),1)
    def test_cached_free_response_hash_and_source_are_checked(self):
        t=self.free_task();source=self.rt.source_for_task({'id':t,'mode':'real'});source.search(1)
        with self.store.transaction() as db:db.execute("UPDATE free_pages SET payload='{}'")
        with self.assertRaises(ConnectionPause):source.search(1)
        self.assertEqual(len(FreeClient.calls),1)
    def test_old_hashed_free_cache_cannot_bypass_strict_pagination(self):
        t=self.free_task();source=self.rt.source_for_task({'id':t,'mode':'real'});source.search(1)
        row=self.store.rows('SELECT * FROM free_pages')[0];payload=json.loads(row['payload'])
        payload['raw']['pagination']['total_entries']=1000
        payload['raw']['pagination']['total_pages']=1
        payload['pagination']=payload['raw']['pagination'];payload['exhausted']=True
        serialized=json.dumps(payload,ensure_ascii=False,sort_keys=True,separators=(',',':'))
        with self.store.transaction() as db:db.execute('UPDATE free_pages SET payload=?,payload_hash=?',(serialized,hashlib.sha256(serialized.encode()).hexdigest()))
        with self.assertRaises(ConnectionPause):source.search(1)
        self.assertEqual(len(FreeClient.calls),1)
    def test_import_needs_no_apollo_key_and_imported_email_is_not_evidence(self):
        t,info=self.imported(target=1)
        self.assertEqual(self.e.counts(t)['合格有邮箱'],0)
        with patch('yingba.runtime.PaidGateway.execute',side_effect=AssertionError('paid gateway')):self.e.run(t)
        self.assertEqual(self.e.counts(t)['合格有邮箱'],2)
        self.assertEqual(FreeClient.calls,[]);self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
        self.assertEqual({r['address'] for r in self.store.rows('SELECT address FROM emails')},{'sales@oil1.example','sales@oil2.example'})
        p,a=export(self.store,t,self.root/'out');wb=load_workbook(p);self.assertEqual(wb['公司总表'].max_row-1,2);wb.close()
        audit=json.loads(a.read_text());self.assertIn(info['sha256'],json.dumps(audit))
    def test_import_exhaustion_can_append_to_same_task_and_keep_surplus(self):
        t,first=self.imported(target=3);self.e.run(t)
        self.assertEqual(self.store.rows('SELECT status FROM tasks WHERE id=?',(t,))[0]['status'],'insufficient')
        _,second=self.imported((3,4),target=3)
        # The second helper task is not run; batch is appended to the original.
        self.assertEqual(append_import(self.e,t,second['batch_id'])['rows'],2)
        self.e.run(t)
        self.assertEqual(self.e.counts(t)['合格有邮箱'],4)
        self.assertEqual(append_import(self.e,t,second['batch_id'])['rows'],0)
        self.assertEqual(self.store.rows('SELECT * FROM charges'),[])
    def test_append_refuses_running_lease_and_non_import_task(self):
        t,info=self.imported();self.e.run(t)
        with self.store.transaction() as db:db.execute('UPDATE tasks SET owner=?,lease=? WHERE id=?',('live',self.e.clock()+600,t))
        with self.assertRaises(ConnectionPause):append_import(self.e,t,info['batch_id'])
        paid=self.e.create(3,'real',config=settings({'estimate_points':1}))
        with self.assertRaises(ValueError):append_import(self.e,paid,info['batch_id'])
    def test_account_and_organization_identity_namespaces_do_not_collide(self):
        from yingba.normalize import keys
        values=dict(keys({'name':'a','domain':'a.example','apollo_id':'same','apollo_account_id':'same'}))
        self.assertEqual(values['apollo'],'same');self.assertEqual(values['apollo_account'],'same')
    def test_source_settings_never_treat_estimated_zero_as_free_paid_call(self):
        with self.assertRaises(ValueError):settings({'source':'apollo_organizations','estimate_points':0})
        self.assertIsNone(settings({'source':'apollo_accounts','estimate_points':50})['estimate_points'])
        with self.assertRaises(ValueError):settings({'source':'apollo_accounts','filters':{'q_organization_keyword_tags':['lubricants']}})

    def batch(self,rows,headers='Company Name,Website,Apollo Account ID,Organization ID'):
        raw=(headers+'\n'+'\n'.join(rows)+'\n').encode()
        return prepare_import(self.store,raw,'candidate.csv','https://app.apollo.io/#/accounts',True)

    def task_for_batch(self,batch,target=3):
        return self.e.create(target,'real',config=settings({'source':'apollo_import','batch_id':batch['batch_id'],'per_page':2}))

    def test_append_initial_batch_is_noop_for_complete_and_insufficient_tasks(self):
        for ids,target in (((1,2),1),((3,4),3)):
            task,batch=self.imported(ids,target);self.e.run(task)
            before=self.store.rows('SELECT * FROM tasks WHERE id=?',(task,))
            audit=self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,))
            events=self.store.rows('SELECT * FROM events WHERE task=?',(task,))
            result=append_import(self.e,task,batch['batch_id'])
            self.assertEqual(result['rows'],0)
            self.assertEqual(self.store.rows('SELECT * FROM tasks WHERE id=?',(task,)),before)
            self.assertEqual(self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,)),audit)
            self.assertEqual(self.store.rows('SELECT * FROM events WHERE task=?',(task,)),events)

    def test_retried_append_before_worker_claim_is_idempotent(self):
        task,_=self.imported(target=3);self.e.run(task)
        batch=self.batch(['Fixture Oil 3,oil3.example,acct-3,org-3','Fixture Oil 4,oil4.example,acct-4,org-4'])
        self.assertEqual(append_import(self.e,task,batch['batch_id'])['rows'],2)
        snapshot=self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,))
        self.assertEqual(append_import(self.e,task,batch['batch_id'])['rows'],0)
        self.assertEqual(self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,)),snapshot)
        self.assertEqual(self.store.rows('SELECT status FROM tasks WHERE id=?',(task,))[0]['status'],'ready')
        self.e.run(task);self.assertEqual(self.e.counts(task)['合格有邮箱'],4)

    def test_concurrent_same_append_commits_exactly_once(self):
        import concurrent.futures
        task,_=self.imported(target=3);self.e.run(task)
        batch=self.batch(['Fixture Oil 3,oil3.example,acct-3,org-3','Fixture Oil 4,oil4.example,acct-4,org-4'])
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(lambda _:append_import(self.e,task,batch['batch_id']),range(2)))
        self.assertEqual(sorted(result['rows'] for result in results),[0,2])
        self.assertEqual(len(self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,))),4)
        self.assertEqual(len(self.store.rows('SELECT * FROM events WHERE task=? AND kind=?',(task,'apollo_import_batch:'+batch['batch_id']))),1)

    def test_append_partial_failure_rolls_back_rows_event_and_task_status(self):
        task,_=self.imported(target=3);self.e.run(task)
        batch=self.batch(['Fixture Oil 3,oil3.example,acct-3,org-3','Fixture Oil 4,oil4.example,acct-4,org-4'])
        original=self.e._ingest;calls=[]
        def fail_on_second(*args):
            calls.append(True)
            if len(calls)==2:raise RuntimeError('synthetic interrupted append')
            return original(*args)
        with patch.object(self.e,'_ingest',side_effect=fail_on_second):
            with self.assertRaises(RuntimeError):append_import(self.e,task,batch['batch_id'])
        self.assertEqual(len(self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,))),2)
        self.assertFalse(self.store.rows('SELECT * FROM events WHERE task=? AND kind=?',(task,'apollo_import_batch:'+batch['batch_id'])))
        self.assertEqual(self.store.rows('SELECT status FROM tasks WHERE id=?',(task,))[0]['status'],'insufficient')
        self.assertEqual(append_import(self.e,task,batch['batch_id'])['rows'],2)
        self.e.run(task);self.assertEqual(self.e.counts(task)['合格有邮箱'],4)

    def test_same_account_with_changed_name_domain_and_org_id_is_one_company(self):
        batch=self.batch(['Fixture Oil 1,oil1.example,STABLE-ACCOUNT,old-org',
                          'Fixture Oil 2,oil2.example,stable-account,new-org'])
        task=self.task_for_batch(batch,target=2);self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],1)
        self.assertEqual(self.store.rows('SELECT status FROM tasks WHERE id=?',(task,))[0]['status'],'insufficient')
        rows=self.store.rows('SELECT outcome FROM raw_records WHERE task=? ORDER BY id',(task,))
        self.assertIn('重复',rows[1]['outcome'])
        self.assertFalse(self.store.rows("SELECT * FROM identities WHERE kind='domain' AND value='oil2.example'"))

    def test_account_dedupes_between_saved_api_and_export_with_new_company_details(self):
        saved=self.free_task(target=1);self.e.run(saved)
        batch=self.batch(['Fixture Oil 3,oil3.example,acct-1,new-org'])
        task=self.task_for_batch(batch,target=1);self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],0)
        self.assertEqual(self.e.counts(saved)['合格有邮箱'],2)
        self.assertIn('重复',self.store.rows('SELECT outcome FROM raw_records WHERE task=?',(task,))[0]['outcome'])

    def test_saved_api_changed_names_domains_across_pages_do_not_meet_target(self):
        task=self.free_task(target=3)
        def response(request):
            offset=(request['page']-1)*2
            rows=[{'id':f'acct-{i}','organization_id':f'org-{offset+i}',
                   'name':f'Fixture Oil {offset+i}','domain':f'oil{offset+i}.example'} for i in (1,2)]
            return parse_accounts_response({'accounts':rows,'pagination':{'page':request['page'],'per_page':2,'total_entries':4,'total_pages':2}},request)
        with patch.object(self.rt._client,'search_accounts',side_effect=response) as calls:self.e.run(task)
        self.assertEqual(calls.call_count,2)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],2)
        self.assertEqual(self.e.counts(task)['距目标差额'],1)
        self.assertEqual(self.store.rows('SELECT status,page FROM tasks WHERE id=?',(task,))[0],{'status':'insufficient','page':3})
        self.assertEqual(len(self.store.rows('SELECT * FROM raw_records WHERE task=?',(task,))),4)

    def test_identical_account_and_org_id_strings_do_not_merge_distinct_companies(self):
        batch=self.batch(['Fixture Oil 1,oil1.example,SHARED-TOKEN,org-1',
                          'Fixture Oil 2,oil2.example,acct-2,SHARED-TOKEN'])
        task=self.task_for_batch(batch,target=2);self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],2)
        identities=self.store.rows("SELECT kind,company FROM identities WHERE value='shared-token' ORDER BY kind")
        self.assertEqual([row['kind'] for row in identities],['apollo','apollo_account'])
        self.assertNotEqual(identities[0]['company'],identities[1]['company'])

    def test_historical_account_id_alone_excludes_renamed_company(self):
        path=self.root/'old.xlsx';wb=Workbook();wb.active.append(['历史保存公司ID']);wb.active.append(['STABLE-ACCOUNT']);wb.save(path)
        self.assertEqual(import_history(self.store,path,{'apollo_account_id':'历史保存公司ID'}),1)
        batch=self.batch(['Fixture Oil 9,oil9.example,stable-account,changed-org'])
        task=self.task_for_batch(batch,target=1);self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],0)
        self.assertTrue(self.store.rows("SELECT * FROM identities WHERE kind='apollo_account' AND value='stable-account' AND historical=1"))
        self.assertFalse(self.store.rows("SELECT * FROM identities WHERE kind='apollo' AND value='stable-account'"))

    def test_export_retains_separate_ids_and_roundtrip_history_dedupes_new_details(self):
        task=self.free_task();self.e.run(task)
        path,_=export(self.store,task,self.root/'out');wb=load_workbook(path,read_only=True)
        rows=list(wb['公司总表'].values);header=list(rows[0]);first=dict(zip(header,rows[1]))
        self.assertEqual(first['Apollo ID'],'org-1');self.assertEqual(first['Apollo Account ID'],'acct-1')
        self.assertEqual(header[:9],['公司ID','公司名称','域名','Apollo ID','邮箱','来源URL','采集时间UTC','核验依据','状态']);wb.close()
        new_store=Store(self.root/'fresh-db');self.assertEqual(import_history(new_store,path),4)
        rt=Runtime(new_store,client_factory=FreeClient,verifier=PublicVerifier(Fetcher()),reviewer=Reviewer(),cost_service=FakeCosts())
        engine=Engine(new_store,source_factory=rt.source_for_task)
        raw=b'Company Name,Website,Apollo Account ID\nFixture Oil 9,oil9.example,acct-1\n'
        batch=prepare_import(new_store,raw,'renamed.csv','https://app.apollo.io/#/accounts',True)
        new_task=engine.create(1,'real',config=settings({'source':'apollo_import','batch_id':batch['batch_id']}));engine.run(new_task)
        self.assertEqual(engine.counts(new_task)['合格有邮箱'],0)

    def test_model_cannot_promote_an_imported_unobserved_email(self):
        task,_=self.imported(target=1)
        def claim_imported(record,evidence,enabled=False):
            result=Reviewer().review_company(record,evidence,enabled)
            result['email_refs']=[{'email':'unverified@'+record['domain'],'evidence_id':'0','quote':'unverified@'+record['domain']}]
            return result
        with patch.object(self.rt.reviewer,'review_company',side_effect=claim_imported):self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],0)
        self.assertEqual(self.store.rows('SELECT * FROM emails'),[])

    def test_model_business_claim_cannot_override_unrelated_public_business(self):
        from yingba.public_web import Page
        class Unrelated:
            def fetch(self,url,expected_domain):
                number=expected_domain.removeprefix('oil').split('.')[0]
                return Page(url,f'<h1>Fixture Oil {number}</h1><p>We sell office chairs.</p><p>sales@{expected_domain}</p>')
        self.rt.verifier=PublicVerifier(Unrelated())
        task,_=self.imported(target=1);self.e.run(task)
        self.assertEqual(self.e.counts(task)['合格有邮箱'],0)
        self.assertEqual(self.e.counts(task)['排除'],2)

    def test_unknown_or_nonzero_cost_decision_does_not_call_free_endpoint(self):
        for decision in ({'classification':'unknown','max_points':None},{'classification':'paid','max_points':1},{'classification':'free','max_points':True}):
            task=self.free_task()
            with patch.object(self.rt.costs,'assess',return_value=decision):self.e.run(task)
            self.assertEqual(FreeClient.calls,[])
            self.assertEqual(self.store.rows('SELECT status FROM tasks WHERE id=?',(task,))[0]['status'],'paused')
            self.assertEqual(self.store.rows('SELECT * FROM charges'),[])

if __name__=='__main__':unittest.main()
