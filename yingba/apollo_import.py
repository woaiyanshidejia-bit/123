"""Read an existing, user-declared Apollo export without calling any service.

Apollo provenance here is a declaration plus an immutable original, not an online
authentication. Imported emails and Apollo's email status never become evidence.
Account IDs and organization IDs remain separate namespaces.
"""
from datetime import date, datetime, timezone
from io import BytesIO, StringIO
import csv
import hashlib
import json
import math
import re
import unicodedata
from urllib.parse import urlsplit, parse_qsl, unquote
import zipfile

from openpyxl import load_workbook

from .normalize import domain

MAX_BYTES = 10_000_000
MAX_ROWS = 100_000
MAX_COLUMNS = 256
ATTESTATION = '用户声明来自已有Apollo导出；未在线认证来源；表中邮箱未经本程序核实'


def _key(value):
    return ''.join(c for c in unicodedata.normalize('NFKC', str(value or '')).casefold() if c.isalnum())


ALIASES = {
    'name': {'companyname', 'accountname', 'organizationname', '公司名称'},
    'domain': {'website', 'companywebsite', 'accountwebsite', 'accountwebsitedomain', 'companydomain', 'domain', 'primarydomain', '域名'},
    'apollo_account_id': {'apolloaccountid', 'accountid'},
    'organization_id': {'organizationid', 'apolloorganizationid'},
}
SENSITIVE = {'apikey', 'apollokey', 'apolloapikey', 'openaiapikey', 'codexapikey', 'password',
             'passwd', 'secret', 'clientsecret', 'authorization', 'bearertoken', 'accesstoken',
             'refreshtoken', 'authtoken', 'token', 'cookies', 'cookie', 'setcookie', 'auth',
             '密钥', '密码', '认证凭据', '访问令牌'}


def _sensitive(key):
    return (key in SENSITIVE or 'apikey' in key or key.startswith('authorization')
            or key.endswith(('password', 'passwordhash', 'accesstoken', 'refreshtoken', 'apitoken', 'clientsecret')))


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def _source_url(value):
    if not isinstance(value, str) or not value.strip() or len(value) > 4096 or any(ord(c) < 32 for c in value):
        raise ValueError('请填写Apollo原始搜索、列表或导出来源链接')
    value = value.strip()
    try:
        url = urlsplit(value)
        host = (url.hostname or '').casefold()
        if (url.scheme != 'https' or url.username is not None or url.password is not None
                or url.port not in (None, 443) or not (host == 'apollo.io' or host.endswith('.apollo.io'))
                or '\\' in value or any(ord(c) < 33 for c in value)):
            raise ValueError
        # Apollo search links may use a hash route. Never retain OAuth callback
        # URLs or credentials pasted in query/hash parameters as provenance.
        parameter_parts = [url.query, url.fragment.split('?', 1)[-1]]
        for part in parameter_parts:
            decoded = unquote(part)
            for key, _ in parse_qsl(decoded.replace(';', '&'), keep_blank_values=True):
                normalized = _key(key)
                if _sensitive(normalized) or normalized in ('key', 'secretkey', 'authkey', 'sessiontoken', 'sessionid', 'code', 'credential', 'credentials'):
                    raise ValueError
    except ValueError:
        raise ValueError('来源须为HTTPS的官方apollo.io链接，不接受凭据、其它域名或端口') from None
    return value


def _value(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError('表中含非有限数字，请修正后导入')
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise ValueError('表中含不支持的单元格值')


def _text(value):
    return '' if value is None else str(value).strip()


def _csv_rows(content):
    try:
        if content.startswith((b'\xff\xfe', b'\xfe\xff')):
            text = content.decode('utf-16', errors='strict')
        else:
            text = content.decode('utf-8-sig', errors='strict')
    except UnicodeError:
        raise ValueError('CSV编码无法识别；请另存为UTF-8或带BOM的UTF-16，不覆盖原件') from None
    if '\x00' in text:
        raise ValueError('CSV含空字符或缺少UTF-16编码标记')
    # Apollo exports use CSV. Supporting a locale's semicolon/tab-separated
    # save is safe; ambiguous sniffing falls back to the standard comma format.
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    try:
        reader = csv.reader(StringIO(text, newline=''), dialect=dialect, strict=True)
        for row in reader:
            yield row
    except csv.Error:
        raise ValueError('CSV格式无效或字段过长，请检查引号与分隔符') from None


def _xlsx_rows(content):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if (len(infos) > 200 or sum(i.file_size for i in infos) > 64_000_000
                    or any(i.flag_bits & 1 for i in infos)):
                raise ValueError('工作簿展开过大或已加密，请拆分后导入')
            if any('vbaproject' in i.filename.casefold() for i in infos):
                raise ValueError('不接受包含宏的文件')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError):
        raise ValueError('不是可读取的XLSX工作簿') from None
    try:
        # Every sheet is part of the retained original bytes. A hidden sheet's
        # obvious credential columns must not escape the import safety gate.
        for candidate in wb:
            first = next(candidate.iter_rows(min_row=1, max_row=1, max_col=MAX_COLUMNS, values_only=True), ())
            if any(_sensitive(_key(value)) for value in first):
                raise ValueError('检测到敏感凭据列（含其它工作表），禁止导入')
        if '导出说明' in wb.sheetnames:
            for row in wb['导出说明'].iter_rows(min_row=1, max_row=20, max_col=2, values_only=True):
                if row[0] == '模式' and ('模拟' in str(row[1]) or 'simulation' in str(row[1]).casefold()):
                    raise ValueError('模拟导出不能作为真实Apollo候选导入')
        sheet = wb.active
        if not sheet or (sheet.max_column and sheet.max_column > MAX_COLUMNS):
            raise ValueError('候选表为空或超过256列')
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ValueError('候选表超过10万行，请拆分后导入')
        yield from sheet.iter_rows(values_only=True)
    finally:
        wb.close()


def _mapping(headers, mapping):
    normalized = [_key(value) for value in headers]
    named = [key for key in normalized if key]
    if len(set(named)) != len(named):
        raise ValueError('候选表含重名列（忽略大小写和空格），请先整理列名')
    if any(_sensitive(key) for key in named):
        raise ValueError('检测到敏感凭据列，禁止将密钥或认证文件导入工作台')
    if mapping is not None and (not isinstance(mapping, dict) or any(k not in ALIASES or not isinstance(v, str) or not v.strip() for k, v in mapping.items())):
        raise ValueError('字段映射仅支持name/domain/apollo_account_id/organization_id')
    result = {}
    for field, aliases in ALIASES.items():
        if mapping is not None and field in mapping:
            key = _key(mapping[field])
            if key not in normalized:
                raise ValueError('字段映射所指的列不存在')
            if field in ('organization_id', 'apollo_account_id') and key not in aliases:
                raise ValueError('账户和组织ID只能映射明确对应列；禁止把联系人或record ID当组织ID')
            result[field] = normalized.index(key)
        else:
            found = [i for i, key in enumerate(normalized) if key in aliases]
            if len(found) > 1:
                raise ValueError('候选列含歧义，请为name/domain等字段明确配置映射')
            if found:
                result[field] = found[0]
    if 'name' not in result or 'domain' not in result:
        raise ValueError('需要公司名称和官网/域名列；不从人名、邮箱或公司名猜测官网')
    if len(set(result.values())) != len(result):
        raise ValueError('不同字段不能映射到同一列')
    return result


def prepare_import(store, content: bytes, filename: str, source_url: str,
                   confirmed: bool, mapping: dict | None = None) -> dict:
    """Persist original bytes and candidate rows; performs no network or inference."""
    if confirmed is not True:
        raise ValueError('请确认这是本人已有的Apollo导出；程序不会代为导出或购买数据')
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_BYTES:
        raise ValueError('候选文件须为非空文件且不超过10MB')
    if not isinstance(filename, str) or not filename.strip() or len(filename) > 255 or any(ord(c) < 32 for c in filename):
        raise ValueError('文件名无效')
    filename = filename.replace('\\', '/').rsplit('/', 1)[-1]
    if filename.upper().startswith('SIMULATION_'):
        raise ValueError('模拟文件不能作为真实Apollo候选导入')
    suffix = filename.rsplit('.', 1)[-1].casefold()
    if suffix not in ('csv', 'xlsx'):
        raise ValueError('仅支持CSV或XLSX；不读取认证文件、XLS或宏工作簿')
    source_url = _source_url(source_url)
    iterator = iter(_csv_rows(content) if suffix == 'csv' else _xlsx_rows(content))
    try:
        header = next(iterator, None)
        if not header or len(header) > MAX_COLUMNS:
            raise ValueError('候选表为空或超过256列')
        headers = [_text(value) for value in header]
        resolved = _mapping(headers, mapping)
        source_hash = hashlib.sha256(content).hexdigest()
        identity = {'file_sha256': source_hash, 'mapping': resolved, 'source_url': source_url,
                    'format': suffix, 'parser_version': 1}
        batch_id = hashlib.sha256(_json(identity).encode('utf-8')).hexdigest()
        imported_at = datetime.now(timezone.utc).isoformat()
        existing = store.rows('SELECT rows,imported_at FROM apollo_imports WHERE batch_id=?', (batch_id,))
        if existing:
            return {'batch_id': batch_id, 'rows': existing[0]['rows'], 'sha256': source_hash,
                    'duplicate': True, 'source_attestation': ATTESTATION, 'imported_at': existing[0]['imported_at']}
        raw_rows = []
        for number, values in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise ValueError('候选表超过10万行，请拆分后导入')
            if len(values) > len(headers) and any(v not in (None, '') for v in values[len(headers):]):
                raise ValueError('数据行比列名多，拒绝错位映射')
            values = [_value(v) for v in values[:len(headers)]]
            values += [None] * (len(headers) - len(values))
            if all(v in (None, '') for v in values):
                continue
            raw = {'headers': headers, 'values': values}
            for i, value in enumerate(values):
                if _key(headers[i]) in ('simulation', 'simulationflag', '模拟标记', '运行模式', '模式') and str(value).strip().casefold() not in ('', 'none', 'false', '0', 'real', '真实', '真实数据'):
                    raise ValueError('检测到模拟标记，不能作为真实Apollo候选导入')
            fields = {key: _text(values[index]) for key, index in resolved.items()}
            record = {'name': fields.get('name', ''), 'domain': domain(fields.get('domain', '')),
                      'apollo_id': fields.get('organization_id', ''),
                      'apollo_account_id': fields.get('apollo_account_id', ''),
                      'website_url': fields.get('domain', ''), 'source_url': source_url,
                      'collected': imported_at, 'apollo_import': {
                          'batch_id': batch_id, 'row_number': number, 'original_file_sha256': source_hash,
                          'source_attestation': ATTESTATION, 'raw': raw}}
            raw_rows.append((batch_id, number, _json(raw), _json(record)))
        if not raw_rows:
            raise ValueError('候选文件没有数据行')
        with store.transaction() as db:
            prior = db.execute('SELECT rows,imported_at FROM apollo_imports WHERE batch_id=?', (batch_id,)).fetchone()
            if prior:
                return {'batch_id': batch_id, 'rows': prior['rows'], 'sha256': source_hash,
                        'duplicate': True, 'source_attestation': ATTESTATION, 'imported_at': prior['imported_at']}
            db.execute('INSERT INTO apollo_imports VALUES(?,?,?,?,?,?,?,?,?)',
                       (batch_id, filename, source_hash, content, source_url, ATTESTATION,
                        imported_at, _json(resolved), len(raw_rows)))
            db.executemany('INSERT INTO apollo_import_rows VALUES(?,?,?,?)', raw_rows)
        return {'batch_id': batch_id, 'rows': len(raw_rows), 'sha256': source_hash, 'duplicate': False,
                'source_attestation': ATTESTATION, 'imported_at': imported_at}
    finally:
        close = getattr(iterator, 'close', None)
        if close:
            close()


def page_import(store, batch_id, page, per_page):
    """Read immutable local candidate pages. Emails remain unreviewed raw data."""
    if not isinstance(batch_id, str) or not re.fullmatch(r'[0-9a-f]{64}', batch_id):
        raise ValueError('候选批次ID无效')
    if type(page) is not int or not 1 <= page <= MAX_ROWS or type(per_page) is not int or not 1 <= per_page <= 1000:
        raise ValueError('本地候选分页参数无效')
    with store.connect() as db:
        db.execute('BEGIN')
        batch = db.execute('SELECT rows FROM apollo_imports WHERE batch_id=?', (batch_id,)).fetchone()
        if not batch:
            raise ValueError('候选批次不存在，请先导入本人已有Apollo文件')
        offset = (page - 1) * per_page
        rows = db.execute('SELECT record FROM apollo_import_rows WHERE batch_id=? ORDER BY row_number LIMIT ? OFFSET ?',
                          (batch_id, per_page, offset)).fetchall()
        return [json.loads(row['record']) for row in rows], offset + len(rows) >= batch['rows']
