import hashlib
import json
import os
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from .engine import stamp,event,qualified_count,QUALIFIED_ROWS
from .normalize import keys,email,cell,source_emails

DEFAULT_MAPPING={'name':'公司名称','domain':'域名','apollo_id':'Apollo ID','apollo_account_id':'Apollo Account ID','email':'邮箱'}
def import_history(store,path,mapping=None):
    """Read only, no writes to historical original. Ambiguous headers fail closed."""
    path=Path(path); mapping=DEFAULT_MAPPING if mapping is None else mapping
    if path.suffix.lower()!='.xlsx':raise ValueError('请先将历史表另存为xlsx（不覆盖原件）')
    if not isinstance(mapping,dict) or not mapping or any(k not in DEFAULT_MAPPING or not isinstance(v,str) or not v.strip() for k,v in mapping.items()) or len(set(mapping.values()))!=len(mapping):raise ValueError('字段映射无效')
    if path.stat().st_size>10_000_000:raise ValueError('历史文件最大10MB')
    content=path.read_bytes()
    if len(content)>10_000_000:raise ValueError('历史文件最大10MB')
    with zipfile.ZipFile(BytesIO(content)) as archive:
        infos=archive.infolist()
        if len(infos)>200 or sum(i.file_size for i in infos)>64_000_000:raise ValueError('历史工作簿展开过大，请分批导入')
    digest=hashlib.sha256(content).hexdigest()
    mapping_key=json.dumps(mapping,ensure_ascii=False,sort_keys=True,separators=(',',':'))
    wb=load_workbook(BytesIO(content),read_only=True,data_only=True,keep_links=False)
    try:
        if '导出说明' in wb.sheetnames:
            for row in wb['导出说明'].iter_rows(min_row=1,max_row=20,max_col=2,values_only=True):
                if row[0]=='模式' and str(row[1]).startswith('模拟'):raise ValueError('模拟导出不能作为真实历史名单导入')
        if wb.active.max_column and wb.active.max_column>256:raise ValueError('历史表超过256列，请整理后导入')
        if wb.active.max_row and wb.active.max_row>100001:raise ValueError('历史表超过10万行，请分批导入')
        rows=wb.active.iter_rows(values_only=True)
        first=next(rows,None)
        if not first:raise ValueError('历史表为空')
        header=list(first)
        named_headers=[value for value in header if value is not None and value!='']
        if len(named_headers)!=len(set(named_headers)):raise ValueError('历史表含重复列名，请调整字段映射')
        indexes={k:header.index(v) for k,v in mapping.items() if v in header}
        if not any(k in indexes for k in DEFAULT_MAPPING):raise ValueError('未找到可去重字段，请配置字段映射')
        count=0
        with store.transaction() as db:
            # Content and field mapping jointly identify an import. A corrected
            # mapping can add missing emails from the same unchanged original.
            if 'mapping' not in {row[1] for row in db.execute('PRAGMA table_info(history_imports)')}:
                db.execute("ALTER TABLE history_imports ADD COLUMN mapping TEXT NOT NULL DEFAULT ''")
            if db.execute('SELECT 1 FROM history_imports WHERE sha256=? AND mapping=?',(digest,mapping_key)).fetchone():return 0
            affected=set()
            for number,row in enumerate(rows,1):
                if number>100000:raise ValueError('历史表超过10万行，请分批导入')
                r={k:row[i] if i<len(row) else None for k,i in indexes.items()}
                identities=keys(r)
                addresses=source_emails(r.get('email',''))
                for address in addresses:identities.append(('email',address))
                if not identities:continue
                for k,v in identities:
                    # Upgrading an existing campaign identity to historical is
                    # meaningful even if import happens while work is paused.
                    for hit in db.execute('SELECT company FROM identities WHERE kind IN (?,?) AND value=? AND company IS NOT NULL',(k,'simulation:'+k,v)):
                        affected.add(hit[0])
                    if k=='email':
                        affected.update(hit[0] for hit in db.execute('SELECT company FROM emails WHERE address=?',(v,)))
                    db.execute('INSERT INTO identities VALUES(?,?,NULL,1) ON CONFLICT(kind,value) DO UPDATE SET historical=1',(k,v))
                count+=1
            tasks=set()
            for cid in affected:
                company=db.execute('SELECT task,stage FROM companies WHERE id=?',(cid,)).fetchone()
                if not company or company['stage']=='excluded':continue
                db.execute("UPDATE companies SET stage='excluded',reason='导入历史名单后发现重复，已排除' WHERE id=?",(cid,))
                db.execute("UPDATE raw_records SET outcome='excluded:历史重复' WHERE company=? AND outcome IN ('candidate','候选','qualified')",(cid,))
                tasks.add(company['task'])
            for task in tasks:
                state=db.execute('SELECT status,target FROM tasks WHERE id=?',(task,)).fetchone()
                if state['status'] in ('complete','simulation_complete'):
                    if qualified_count(db,task)<state['target']:db.execute("UPDATE tasks SET status='insufficient',reason='历史导入剔除重复后未达目标' WHERE id=?",(task,))
                elif state['status'] in ('ready','running'):
                    db.execute("UPDATE tasks SET status='paused',reason='历史名单已更新，请检查排除结果后继续' WHERE id=?",(task,))
                event(db,task,'history_dedup','历史名单更新，重新核对任务有效公司数量')
            db.execute('INSERT INTO history_imports(source,sha256,imported,rows,mapping) VALUES(?,?,?,?,?)',(path.name,digest,stamp(),count,mapping_key))
        return count
    finally:wb.close()

def export(store,task,directory):
    """Single consistent read transaction; official sheets contain qualified only."""
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    with store.connect() as db:
        db.execute('BEGIN')
        found=db.execute('SELECT * FROM tasks WHERE id=?',(task,)).fetchone()
        if not found:raise ValueError('任务不存在')
        t=dict(found)
        rows=[dict(r) for r in db.execute(QUALIFIED_ROWS,(task,))]
        raw=[dict(r) for r in db.execute('SELECT * FROM raw_records WHERE task=?',(task,))]
        evidence=[dict(r) for r in db.execute('SELECT e.* FROM evidence e JOIN companies c ON e.company=c.id WHERE c.task=?',(task,))]
        pending_evidence=[dict(r) for r in db.execute('SELECT * FROM verification_cache WHERE task_id=?',(task,))]
        excluded=[dict(r) for r in db.execute("SELECT id,name,domain,stage,reason FROM companies WHERE task=? AND stage!='qualified'",(task,))]
    grouped={}
    for r in rows:grouped.setdefault(r['id'],[]).append(r)
    wb=Workbook(); companies=wb.active;companies.title='公司总表'
    # Preserve existing column positions; account IDs get their own column and
    # never replace organization IDs in the historical Apollo ID column.
    companies.append(['公司ID','公司名称','域名','Apollo ID','邮箱','来源URL','采集时间UTC','核验依据','状态','Apollo Account ID'])
    for cid,rs in grouped.items():
        r=rs[0];record=json.loads(r['raw']);account=record.get('apollo_account_id') or record.get('account_id') or ''
        companies.append([cid,cell(r['name']),cell(r['domain']),cell(r['apollo']),cell('; '.join(x['address'] for x in rs)),cell('\n'.join(dict.fromkeys(x['url'] for x in rs))),cell('\n'.join(dict.fromkeys(x['evidence_time'] for x in rs))),cell('\n'.join(dict.fromkeys(x['basis'] for x in rs))),'模拟合格' if t['mode']=='simulation' else '公开核实；非投递验证',cell(account)])
    sheet=wb.create_sheet('唯一邮箱清单');sheet.append(['邮箱','归属公司ID','公司名称','来源URL','采集时间UTC','核验依据'])
    by_address={}
    for r in rows:by_address.setdefault(r['address'],[]).append(r)
    for address,rs in by_address.items():sheet.append([cell(address),'; '.join(str(r['id']) for r in rs),cell('; '.join(r['name'] for r in rs)),cell('\n'.join(dict.fromkeys(r['url'] for r in rs))),cell('\n'.join(dict.fromkeys(r['evidence_time'] for r in rs))),cell('\n'.join(dict.fromkeys(r['basis'] for r in rs)))])
    meta=wb.create_sheet('导出说明');meta.append(['项目','值'])
    for k,v in {'模式':'模拟数据，禁止作为真实客户使用' if t['mode']=='simulation' else '真实数据','任务ID':task,'目标':t['target'],'合格公司':len(grouped),'唯一邮箱':len(by_address),'任务状态':t['status'],'说明':'公开来源核实不等于实际投递验证；不含邮件内容；未审核不计达标','ID说明':'Apollo ID为组织ID；Apollo Account ID为已保存公司账户ID，两种ID独立去重','导出UTC':stamp()}.items():meta.append([k,v])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='174A44')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col)+2))
    prefix='SIMULATION_' if t['mode']=='simulation' else 'YINGBA_'
    # Unique artifacts; no silent overwrite of earlier exports.
    token=__import__('uuid').uuid4().hex[:10]
    path=directory/f'{prefix}{task[:8]}_{token}.xlsx'
    audit=path.with_suffix('.audit.json')
    with tempfile.TemporaryDirectory(prefix='.export-',dir=directory) as temp:
        staged=Path(temp)/path.name
        wb.save(staged)
        check=load_workbook(staged,read_only=True)
        try:
            if check['公司总表'].max_row-1!=len(grouped) or check['唯一邮箱清单'].max_row-1!=len(by_address):raise ValueError('导出回读对账失败')
            for expected in wb:
                actual=check[expected.title]
                expected_values=[tuple(c.value for c in row) for row in expected]
                actual_values=[tuple(c.value for c in row) for row in actual]
                # openpyxl represents an exported empty string as None.
                normalize=lambda rows:[tuple(None if value=='' else value for value in row) for row in rows]
                if normalize(expected_values)!=normalize(actual_values):raise ValueError('导出逐格回读对账失败')
                if any(c.data_type=='f' for row in actual for c in row):raise ValueError('导出检测到公式，已阻止发布')
        finally:check.close()
        staged_audit=Path(temp)/audit.name
        staged_audit.write_text(json.dumps({'task':t,'simulation':t['mode']=='simulation','raw_records':raw,'evidence':evidence,'fetched_verification_cache':pending_evidence,'excluded':excluded,'counts':{'companies':len(grouped),'unique_emails':len(by_address)},'xlsx_sha256':hashlib.sha256(staged.read_bytes()).hexdigest()},ensure_ascii=False,indent=2),encoding='utf-8')
        # Publish audit first; Excel is only visible after both files are valid.
        os.replace(staged_audit,audit)
        try:os.replace(staged,path)
        except BaseException:
            audit.unlink(missing_ok=True)
            raise
    return path,audit
